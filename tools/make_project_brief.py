"""Generate the ISCF project-workshop brief for the CSLAP robustness project.

Mirrors the layout of last year's "Fiche projet" (header, title, specialty
tick-boxes, Context, Objectives, Evaluation) so the two briefs sit together in
the same series. Produces an ordinary .docx that can be edited freely in Word.

Run:  python tools/make_project_brief.py
Out:  ISCF_project_CSLAP_robustness.docx  (repo root)
"""
from __future__ import annotations

import os

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_TAB_ALIGNMENT
from docx.shared import Pt, RGBColor, Cm

REPO = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
OUT = os.path.join(REPO, "ISCF_project_CSLAP_robustness.docx")

TICKED = "☒"      # box with X
EMPTY = "☐"       # empty box


# --------------------------------------------------------------- helpers
def rich(par, chunks):
    """Add runs to `par` from (text, bold, italic) triples."""
    for text, bold, italic in chunks:
        r = par.add_run(text)
        r.bold = bold
        r.italic = italic
    return par


def label(doc, text):
    """Underlined section label, as in the reference brief."""
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(12)
    p.paragraph_format.space_after = Pt(6)
    r = p.add_run(text)
    r.underline = True
    return p


def body(doc, chunks, space_after=8):
    p = doc.add_paragraph()
    p.paragraph_format.space_after = Pt(space_after)
    p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    return rich(p, chunks)


def bullet(doc, chunks):
    p = doc.add_paragraph(style="List Bullet")
    p.paragraph_format.space_after = Pt(4)
    return rich(p, chunks)


def numbered(doc, chunks):
    p = doc.add_paragraph(style="List Number")
    p.paragraph_format.space_after = Pt(4)
    return rich(p, chunks)


def formula(doc, text):
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_before = Pt(8)
    p.paragraph_format.space_after = Pt(8)
    r = p.add_run(text)
    r.bold = True
    r.font.size = Pt(12)
    return p


def placeholder(doc, text):
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_before = Pt(6)
    p.paragraph_format.space_after = Pt(10)
    r = p.add_run(text)
    r.italic = True
    r.font.color.rgb = RGBColor(0x80, 0x80, 0x80)
    r.font.size = Pt(9)
    return p


# --------------------------------------------------------------- document
doc = Document()

sec = doc.sections[0]
sec.top_margin = Cm(2.2)
sec.bottom_margin = Cm(2.2)
sec.left_margin = Cm(2.5)
sec.right_margin = Cm(2.5)

style = doc.styles["Normal"]
style.font.name = "Calibri"
style.font.size = Pt(11)

# running header: "Master ISC" left, "ISCF : Project workshop" right
hp = sec.header.paragraphs[0]
hp.text = ""
tabs = hp.paragraph_format.tab_stops
tabs.add_tab_stop(sec.page_width - sec.left_margin - sec.right_margin,
                  WD_TAB_ALIGNMENT.RIGHT)
hr = hp.add_run("Master ISC\tISCF : Project workshop")
hr.font.size = Pt(9)

# ---------------------------------------------------------------- title
t = doc.add_paragraph()
t.alignment = WD_ALIGN_PARAGRAPH.CENTER
t.paragraph_format.space_after = Pt(18)
tr = t.add_run("Robust Storage Assignment for Warehouse Picking Stations")
tr.italic = True
tr.bold = True
tr.font.size = Pt(20)

st = doc.add_paragraph()
st.alignment = WD_ALIGN_PARAGRAPH.CENTER
st.paragraph_format.space_after = Pt(20)
sr = st.add_run("Designing a workload contract that survives demand drift")
sr.italic = True
sr.font.size = Pt(13)

label(doc, "Project leader:")
doc.add_paragraph()

label(doc, "Specialty(ies) – Tick one (or more) box(es):")
for mark, name in ((TICKED, "AOS Specialty"), (TICKED, "ARS Specialty"),
                   (EMPTY, "BMI Specialty"), (EMPTY, "SMC Specialty"),
                   (EMPTY, "SMT Specialty")):
    p = doc.add_paragraph()
    p.paragraph_format.left_indent = Cm(1.0)
    p.paragraph_format.space_after = Pt(2)
    p.add_run(f"{mark}\t{name}")

label(doc, "Project description:")

# ---------------------------------------------------------------- context
label(doc, "Context :")

placeholder(doc, "[INSERT FIGURE: warehouse 3D overview — reuse the image "
                 "from last year's brief]")

body(doc, [
    ("A ", 0, 0), ("warehouse", 1, 0),
    (" is a distribution center whose core activity is to ", 0, 0),
    ("store and deliver products", 1, 0),
    (". It receives ", 0, 0), ("large pallets of a single product reference", 1, 0),
    (" and outputs ", 0, 0),
    ("customer orders that contain smaller quantities of multiple references", 1, 0),
    (". As in last year's project, we focus only on the ", 0, 0),
    ("picking sector", 1, 0), (" of the warehouse.", 0, 0),
])

body(doc, [
    ("The picking area consists of multiple ", 0, 0), ("aisles", 1, 0),
    (". At the end of each aisle sits a ", 0, 0),
    ("handover location (“station”)", 1, 0),
    (" where items picked from that aisle are consolidated and sent onward for "
     "packing and shipping. A customer order is a small basket of several "
     "product references. To fulfil an order, a tote travels along the line and "
     "must ", 0, 0),
    ("stop at every station that holds at least one of the ordered products", 1, 0),
    (". Each stop is a ", 0, 0), ("station visit", 1, 0),
    (", and visits cost time, buffer space and throughput.", 0, 0),
])

placeholder(doc, "[INSERT FIGURE: top view of the picking sector with aisles, "
                 "handover locations and depot — reuse from last year's brief]")

body(doc, [
    ("This year the project focuses on a single placement lens:", 0, 0),
])
bullet(doc, [
    ("CSLAP — Correlated Storage Location Assignment Problem.", 1, 0),
    (" Assign each SKU to exactly one station so as to minimise ", 0, 0),
    ("order–station visits", 1, 0),
    (", by co-locating frequently co-ordered SKUs within the same station and "
     "thereby reducing the number of handovers per order.", 0, 0),
])

body(doc, [
    ("Two forces pull against each other, and the whole project lives in that "
     "tension. The ", 0, 0), ("objective", 1, 0),
    (" pulls towards ", 0, 0), ("concentration", 1, 0),
    (": merging popular, frequently co-ordered SKUs onto a few stations removes "
     "visits. The ", 0, 0), ("workload constraint", 1, 0),
    (" pushes towards ", 0, 0), ("dispersion", 1, 0),
    (": every station has a bounded picking throughput, and a station that "
     "receives too much work goes into overtime, spills work to the next "
     "period, or saturates its buffer and stalls the line for every order "
     "behind it. Stations are also ", 0, 0),
    ("physically heterogeneous", 1, 0),
    (" — static shelving, flow racks and pallet positions do not have the "
     "same throughput — so an equal split of the work is not the correct "
     "target.", 0, 0),
])

# ------------------------------------------------------- the robustness problem
label(doc, "The problem this project addresses:")

body(doc, [
    ("A storage assignment is computed on ", 0, 0), ("historical orders", 1, 0),
    (" but operated on ", 0, 0), ("future orders", 1, 0),
    (". Demand ", 0, 0), ("drifts", 1, 0),
    (": some references become more popular, others fade, and the composition "
     "of the order stream changes. A layout that respected every station's "
     "workload limit on the data it was built from may therefore ", 0, 0),
    ("overload a station once it is in operation", 1, 0),
    (". The central question of this project is:", 0, 0),
])

q = doc.add_paragraph()
q.paragraph_format.left_indent = Cm(1.0)
q.paragraph_format.right_indent = Cm(1.0)
q.paragraph_format.space_before = Pt(6)
q.paragraph_format.space_after = Pt(10)
qr = q.add_run("How should the workload constraint be written so that a layout "
               "built on the past remains workload-feasible on the future — "
               "and how much does that protection cost in station visits?")
qr.italic = True

body(doc, [
    ("The starting observation is that ", 0, 0),
    ("two different things change between the past and the future", 1, 0),
    (", and they deserve different treatment:", 0, 0),
])
numbered(doc, [
    ("The scale.", 1, 0),
    (" The total amount of work may grow or shrink — a window with 10,000 "
     "picking lines may be followed by one with 100,000. ", 0, 0),
    ("Working hypothesis:", 1, 0),
    (" the warehouse can absorb scale, because more volume is met with more "
     "shifts and more staff.", 0, 0),
])
numbered(doc, [
    ("The mix.", 1, 0),
    (" The way that work distributes across stations may change. ", 0, 0),
    ("Working hypothesis:", 1, 0),
    (" the warehouse cannot absorb an unbalanced mix, because a station's "
     "throughput is fixed by its physical design.", 0, 0),
])

body(doc, [
    ("If those two hypotheses hold, then feasibility should not be expressed as "
     "an ", 0, 0), ("absolute ceiling in lines per station", 1, 0),
    (" (which becomes meaningless as soon as the volume changes), but as a ", 0, 0),
    ("share of whatever total work arrives", 1, 0),
    (", with a tolerance around the share observed historically. That is the "
     "contract this project is asked to build, study and challenge:", 0, 0),
])

formula(doc, "Wₛ(window)  ≤  ( μₛ + z · σₛ ) "
             "× L(window)")

body(doc, [("where", 0, 0)], space_after=4)
for sym, txt in (
    ("Wₛ(window)",
     "the picking lines executed at station s over the window considered;"),
    ("L(window)",
     "the total picking lines executed warehouse-wide over the same window;"),
    ("μₛ",
     "station s's mean share of the total work, measured on the historical "
     "(training) window;"),
    ("σₛ",
     "the variability of that share over the historical window;"),
    ("z",
     "a tolerance multiplier — how far a station is allowed to drift from "
     "its historical share before the layout is judged infeasible."),
):
    p = doc.add_paragraph(style="List Bullet")
    p.paragraph_format.space_after = Pt(3)
    rich(p, [(sym, 1, 0), ("  —  " + txt, 0, 0)])

body(doc, [
    ("Because both sides of the inequality scale with the window's total work, "
     "a pure change of scale cancels out and cannot break the constraint. Only "
     "a change of ", 0, 0), ("mix", 1, 0),
    (" can. This is the property the project must both exploit and verify.", 0, 0),
])

# ---------------------------------------------------------------- given
label(doc, "For this project it will be given:")
bullet(doc, [
    ("A common dataset", 1, 0),
    (": historical customer orders with a chronological ordering, SKU–"
     "station membership, station capacities and throughput speeds.", 0, 0),
])
bullet(doc, [
    ("A reference CSLAP solution", 1, 0),
    (" computed on the same dataset (the “nominal” layout, minimising "
     "station visits without any protection), to serve as the baseline every "
     "robust variant is compared against.", 0, 0),
])
bullet(doc, [
    ("An evaluation harness", 1, 0),
    (" that replays a layout on a held-out order stream and reports station "
     "visits and per-station workload.", 0, 0),
])

# ---------------------------------------------------------------- objectives
label(doc, "Objectives:")

body(doc, [
    ("Baseline and the naive contract.", 1, 0),
    (" Reproduce the nominal CSLAP assignment and evaluate it out of sample "
     "under the simplest possible workload rule: a fixed ceiling per station, "
     "calibrated on the training window (for instance an equal share of the "
     "training work plus a fixed safety margin). Split the order stream ", 0, 0),
    ("chronologically", 1, 0),
    (" — never randomly — at several train/test fractions (50/50, "
     "60/40, 70/30, 80/20, 90/10). Report where this rule gives a verdict you "
     "can trust and where it does not. ", 0, 0),
    ("Pay particular attention to what happens to the verdict as the test "
     "window shrinks.", 1, 0),
])

body(doc, [
    ("Contract design: the share band.", 1, 0),
    (" Implement the volume-normalised contract given above and contrast it "
     "with the fixed ceiling. The essential modelling question is how to obtain "
     "μₛ and σₛ from a single historical window. A mean "
     "share is immediate; a ", 0, 0), ("variance is not", 1, 0),
    (", because one window yields only one observation per station. Students "
     "must design a ", 0, 0), ("segmentation", 1, 0),
    (" of the training window (for example into consecutive blocks of orders) "
     "and justify how many blocks to use — too few gives an unreliable "
     "variance estimate, too many makes each block noisy.", 0, 0),
])

body(doc, [
    ("Investigate the two premises.", 1, 0),
    (" The contract rests on empirical claims that must be tested, not "
     "assumed:", 0, 0),
])
bullet(doc, [
    ("Is a station's share of the work genuinely ", 0, 0),
    ("independent of the total volume", 1, 0),
    ("? If busy periods have a systematically different mix, the contract is "
     "mis-specified exactly when it matters most.", 0, 0),
])
bullet(doc, [
    ("Does the observed variability σₛ reflect ", 0, 0),
    ("real composition drift", 1, 0), (", or merely ", 0, 0),
    ("sampling noise", 1, 0),
    (" arising from the finite size of each block? These behave very "
     "differently: sampling noise averages away as the window grows, drift does "
     "not. A tolerance fitted on small blocks and applied to a large window may "
     "be far too permissive. Propose and implement a way to separate the two.", 0, 0),
])

body(doc, [
    ("Calibration of the tolerance z.", 1, 0),
    (" z is the contract, not a free parameter to be tuned until the results "
     "look good. Determine a defensible range from the data itself: report the "
     "", 0, 0), ("coverage curve", 1, 0),
    (" — for each candidate z, the proportion of historical periods in "
     "which every station stays inside its band. Identify an upper bound (above "
     "which the constraint no longer prevents anything) and a lower bound "
     "(below which the constraint demands a balance the site has never "
     "achieved). Report z as ", 0, 0), ("measured coverage", 1, 0),
    (", not as a number of standard deviations, unless the distribution is "
     "shown to justify it.", 0, 0),
])

body(doc, [
    ("Optimisation under the contract.", 1, 0),
    (" Compute assignments that minimise order–station visits subject to "
     "the share band, for each z in the calibrated range. Arrive at them using "
     "either an ", 0, 0), ("exact optimisation route", 1, 0),
    (" (a MILP formulation solved with a standard solver) or a ", 0, 0),
    ("heuristic/metaheuristic route", 1, 0),
    (" (constructive rules with local search, tabu/ILS, or co-occurrence-aware "
     "swaps). Implementing both is ", 0, 0), ("not", 1, 1),
    (" required. Report the ", 0, 0), ("price of the contract", 1, 0),
    (": the extra station visits incurred, relative to the unprotected "
     "layout.", 0, 0),
])

body(doc, [
    ("A fair competitor (strongly recommended).", 1, 0),
    (" Any protection can be bought more crudely by simply tightening every "
     "station's allowance by a uniform factor. Implement that alternative and "
     "compare it against the share band ", 0, 0), ("at equal cost", 1, 0),
    (" — that is, calibrate the uniform tightening so that it spends the "
     "same number of extra visits, then ask which of the two delivers more "
     "out-of-sample protection. A method is only worth its complexity if it "
     "beats the simple alternative on a level playing field.", 0, 0),
])

body(doc, [
    ("Advanced extension (optional): budgeted uncertainty.", 1, 0),
    (" Instead of a tolerance on the aggregate share, protection can be "
     "targeted at the individual references most likely to deviate, using a ", 0, 0),
    ("Bertsimas–Sim budget of uncertainty", 1, 0),
    (": assume that at most Γ references deviate adversely at once, and "
     "add the corresponding worst-case term to each station's constraint. The "
     "resulting non-linear term admits an exact linear reformulation by "
     "linear-programming duality. Groups taking this route should report what "
     "Γ buys over the share band alone, and should check whether the "
     "per-reference deviations being insured are genuine drift or simply "
     "counting noise.", 0, 0),
])

# --------------------------------------------------------- evaluation
label(doc, "Evaluation & sensitivity:")

body(doc, [
    ("Report, for every train/test fraction and every z considered: the number "
     "of station visits (in and out of sample), the number of stations "
     "exceeding their allowance on the held-out stream, the severity of the "
     "worst station, and the share of total work that spills above the "
     "allowances. Compare each protected variant against both the nominal "
     "layout and the equal-cost uniform tightening.", 0, 0),
])

body(doc, [
    ("Examine sensitivities: the number of blocks used to estimate the "
     "variance, the train/test fraction, the tolerance z, station capacities, "
     "and demand scaling. State clearly which conclusions are stable across "
     "these choices and which are not.", 0, 0),
])

body(doc, [
    ("Two points of scientific method will be weighted heavily.", 1, 0),
    (" First, ", 0, 0), ("the size of the evidence", 1, 0),
    (": a held-out window contains a limited number of independent periods, so "
     "differences of one or two periods between two variants may not be "
     "distinguishable from noise — say so, and quantify it. Second, ", 0, 0),
    ("solver variability", 1, 0),
    (": if a metaheuristic is used, re-run it with several random seeds and "
     "report the spread, so that a difference between two methods can be "
     "separated from a difference between two runs of the same method. ", 0, 0),
    ("A negative or inconclusive result, correctly established, is a valid and "
     "valuable outcome of this project.", 1, 0),
])

body(doc, [
    ("Emphasis throughout is on ", 0, 0), ("reproducibility and clarity", 1, 0),
    (": every reported number should be regenerable from the submitted code "
     "with a single documented command.", 0, 0),
])

doc.save(OUT)
print(f"written: {OUT}")
print(f"size: {os.path.getsize(OUT) / 1024:.1f} KB")
